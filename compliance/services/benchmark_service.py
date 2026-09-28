"""
Servicio para importar benchmarks desde archivos XLSX
"""
import logging
import openpyxl
from decimal import Decimal
from typing import Dict, List, Tuple
from django.db import transaction
from compliance.models import Benchmark, BenchmarkPolicy
from compliance.validators import validate_xlsx_file, SystemConstants

logger = logging.getLogger('compliance')


class BenchmarkImportService:
    """Servicio para importar benchmarks CIS desde XLSX"""
    
    # Mapeo de columnas XLSX a campos del modelo (case-insensitive)
    COLUMN_MAPPING = {
        'section': 'section',
        'recommendation': 'recommendation',
        'profile': 'profile',
        'title': 'title',
        'assessment status': 'assessment_status',
        'description': 'description',
        'rationale statement': 'rationale_statement',
        'impact statement': 'impact_statement',
        'remediation procedure': 'remediation_procedure',
        'audit procedure': 'audit_procedure',
        'additional information': 'additional_information',
        'cis controls': 'cis_controls',
        'cis safeguards 1 v8': 'cis_safeguards_1_v8',
        'cis safeguards 2 v8': 'cis_safeguards_2_v8',
        'cis safeguards 3 v8': 'cis_safeguards_3_v8',
        'v8 ig1': 'v8_ig1',
        'v8 ig2': 'v8_ig2',
        'v8 ig3': 'v8_ig3',
        'cis safeguards 1 v7': 'cis_safeguards_1_v7',
        'cis safeguards 2 v7': 'cis_safeguards_2_v7',
        'cis safeguards 3 v7': 'cis_safeguards_3_v7',
        'v7 ig1': 'v7_ig1',
        'v7 ig2': 'v7_ig2',
        'v7 ig3': 'v7_ig3',
        'references': 'references',
        'default value': 'default_value',
    }
    
    # Profiles válidos (solo Server)
    VALID_PROFILES = {
        'Level 1 - Server': 'L1_SERVER',
        'Level 2 - Server': 'L2_SERVER',
    }
    
    def __init__(self):
        self.errors = []
        self.warnings = []
        self.stats = {
            'total_rows': 0,
            'imported': 0,
            'skipped': 0,
            'errors': 0
        }
    
    def import_benchmark(self, file, benchmark_name: str, version: str, uploaded_by: str = '') -> Tuple[Benchmark, Dict]:
        """
        Importa un benchmark desde archivo XLSX
        
        Returns:
            Tuple[Benchmark, Dict]: (benchmark_object, stats_dict)
        """
        try:
            # Validar archivo
            validate_xlsx_file(file)
            
            # Crear el Benchmark
            benchmark = Benchmark.objects.create(
                name=benchmark_name,
                version=version,
                file_original=file,
                uploaded_by=uploaded_by
            )
            
            logger.info(f"Benchmark creado: {benchmark}")
            logger.info(f"=== LLAMANDO A _process_xlsx ===")
            self._process_xlsx(file, benchmark)
            logger.info(f"=== _process_xlsx COMPLETADO ===")
            
            # Procesar el XLSX
            # self._process_xlsx(file, benchmark)
            
            return benchmark, self.stats
            
        except Exception as e:
            logger.error(f"Error al importar benchmark: {str(e)}")
            self.errors.append(f"Error general: {str(e)}")
            # Si hay error, eliminar el benchmark creado
            if 'benchmark' in locals():
                benchmark.delete()
            raise
    
    @transaction.atomic
    def _process_xlsx(self, file, benchmark: Benchmark):
        """Procesa el archivo XLSX y crea las políticas"""
        
        # Cargar workbook
        wb = openpyxl.load_workbook(file, data_only=True)
        ws = wb.active
        
        # Obtener headers (primera fila) y normalizar a lowercase
        headers = []
        for cell in ws[1]:
            header = str(cell.value).strip().lower() if cell.value else ''
            headers.append(header)
        
        logger.info(f"Headers encontrados (normalizados): {headers}")
        
        # Validar que tenemos las columnas mínimas requeridas
        required_columns = ['section', 'profile', 'title']
        missing_columns = [col for col in required_columns if col not in headers]
        
        if missing_columns:
            raise ValueError(f"Columnas requeridas faltantes: {missing_columns}")
        
        # Crear mapeo de índices
        column_indices = {header: idx for idx, header in enumerate(headers)}
        
        # Procesar filas (desde la fila 2)
        for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            self.stats['total_rows'] += 1
            
            try:
                # PRIMERO: Validar el profile ANTES de extraer todos los datos
                profile_idx = column_indices.get('profile')
                if profile_idx is not None and profile_idx < len(row):
                    profile_raw = str(row[profile_idx]).strip() if row[profile_idx] else ''
                    
                    # Validar que sea un profile Server válido
                    if profile_raw not in self.VALID_PROFILES:
                        self.warnings.append(f"Fila {row_num}: Profile '{profile_raw}' no es Server, se omite")
                        self.stats['skipped'] += 1
                        continue
                else:
                    self.warnings.append(f"Fila {row_num}: Sin columna Profile, se omite")
                    self.stats['skipped'] += 1
                    continue
                
                # SEGUNDO: Extraer todos los datos de la fila (ya sabemos que profile es válido)
                row_data = self._extract_row_data(row, headers, column_indices)
                
                # Validar que tengamos datos mínimos
                if not row_data.get('title'):
                    self.warnings.append(f"Fila {row_num}: Sin título, se omite")
                    self.stats['skipped'] += 1
                    continue
                
                # Crear la política
                self._create_policy(benchmark, row_data, row_num)
                self.stats['imported'] += 1
                
            except Exception as e:
                logger.error(f"Error en fila {row_num}: {str(e)}")
                self.errors.append(f"Fila {row_num}: {str(e)}")
                self.stats['errors'] += 1
        
        logger.info(f"Importación completada: {self.stats}")
    
    def _extract_row_data(self, row: tuple, headers: list, column_indices: dict) -> dict:
        """Extrae datos de una fila del XLSX"""
        
        data = {}
        
        for xlsx_column, model_field in self.COLUMN_MAPPING.items():
            if xlsx_column in column_indices:
                idx = column_indices[xlsx_column]
                value = row[idx] if idx < len(row) else None
                
                # Convertir valor según el tipo de campo
                if model_field in ['v7_ig1', 'v7_ig2', 'v7_ig3', 'v8_ig1', 'v8_ig2', 'v8_ig3']:
                    # Campos booleanos
                    data[model_field] = self._parse_boolean(value)
                elif model_field == 'profile':
                    # Convertir profile a código
                    data[model_field] = self.VALID_PROFILES.get(str(value).strip(), '')
                else:
                    # Campos de texto
                    data[model_field] = str(value).strip() if value else ''
        
        return data
    
    def _parse_boolean(self, value) -> bool:
        """Convierte valores de Excel a booleano"""
        if value is None or value == '':
            return False
        
        if isinstance(value, bool):
            return value
        
        value_str = str(value).strip().lower()
        return value_str in ['yes', 'y', 'true', '1', 'x', '✓', 'si', 'sí']
    
    def _create_policy(self, benchmark: Benchmark, data: dict, row_num: int):
        """Crea una BenchmarkPolicy desde los datos extraídos"""
        
        try:
            policy = BenchmarkPolicy.objects.create(
                benchmark=benchmark,
                section=data.get('section', ''),
                recommendation=data.get('recommendation', ''),
                profile=data.get('profile', ''),
                title=data.get('title', ''),
                assessment_status=data.get('assessment_status', ''),
                description=data.get('description', ''),
                rationale_statement=data.get('rationale_statement', ''),
                impact_statement=data.get('impact_statement', ''),
                remediation_procedure=data.get('remediation_procedure', ''),
                audit_procedure=data.get('audit_procedure', ''),
                additional_information=data.get('additional_information', ''),
                cis_controls=data.get('cis_controls', ''),
                cis_safeguards_1_v8=data.get('cis_safeguards_1_v8', ''),
                cis_safeguards_2_v8=data.get('cis_safeguards_2_v8', ''),
                cis_safeguards_3_v8=data.get('cis_safeguards_3_v8', ''),
                v8_ig1=data.get('v8_ig1', False),
                v8_ig2=data.get('v8_ig2', False),
                v8_ig3=data.get('v8_ig3', False),
                cis_safeguards_1_v7=data.get('cis_safeguards_1_v7', ''),
                cis_safeguards_2_v7=data.get('cis_safeguards_2_v7', ''),
                cis_safeguards_3_v7=data.get('cis_safeguards_3_v7', ''),
                v7_ig1=data.get('v7_ig1', False),
                v7_ig2=data.get('v7_ig2', False),
                v7_ig3=data.get('v7_ig3', False),
                references=data.get('references', ''),
                default_value=data.get('default_value', ''),
            )
            
            logger.debug(f"Política creada: {policy.section} - {policy.title[:50]}")
            return policy
            
        except Exception as e:
            raise Exception(f"Error al crear política: {str(e)}")
    
    def get_report(self) -> dict:
        """Retorna un reporte de la importación"""
        return {
            'success': self.stats['errors'] == 0,
            'stats': self.stats,
            'errors': self.errors,
            'warnings': self.warnings
        }


def auto_version_benchmark(benchmark_name: str) -> str:
    """
    Genera automáticamente la siguiente versión del benchmark
    
    Ejemplo:
        Si existe v1.0.0, retorna v2.0.0
        Si no existe ninguna, retorna v1.0.0
    """
    existing_versions = Benchmark.objects.filter(
        name=benchmark_name
    ).order_by('-version')
    
    if not existing_versions.exists():
        return '1.0.0'
    
    # Obtener última versión
    last_version = existing_versions.first().version
    
    try:
        # Parsear versión (formato: X.Y.Z)
        major, minor, patch = map(int, last_version.split('.'))
        # Incrementar versión mayor
        return f"{major + 1}.0.0"
    except:
        # Si no se puede parsear, incrementar como número simple
        try:
            return str(int(float(last_version)) + 1) + '.0.0'
        except:
            return '1.0.0'