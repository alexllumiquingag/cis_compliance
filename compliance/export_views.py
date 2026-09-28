"""
Vistas para exportación de datos a Excel
"""
from django.shortcuts import render
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.db.models import Q, Max, Prefetch
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from compliance.models import Asset, AssetLayer, Scan, Baseline


@login_required
def export_assets_to_xlsx(request):
    """
    Exporta la lista de assets a Excel respetando los filtros aplicados
    """
    
    # ============================================================
    # OBTENER LOS MISMOS FILTROS QUE LA VISTA DE LISTA
    # ============================================================
    
    # Filtros básicos
    status_filter = request.GET.get('status', 'ACTIVE')
    search_query = request.GET.get('search', '').strip()
    
    # Filtros de clasificación
    environment_filter = request.GET.get('environment', '')
    administrator_filter = request.GET.get('administrator', '')
    responsible_filter = request.GET.get('responsible', '')
    pci_filter = request.GET.get('pci_standard', '')
    
    # Filtros de capa y compliance
    layer_type_filter = request.GET.get('layer_type', '')
    compliance_min = request.GET.get('compliance_min', '')
    compliance_max = request.GET.get('compliance_max', '')
    
    # Filtro de fecha de último escaneo
    scan_date_from = request.GET.get('scan_date_from', '')
    scan_date_to = request.GET.get('scan_date_to', '')
    
    # Filtro de cambio en compliance
    show_only_changed = request.GET.get('show_only_changed', '') == 'on'
    
    # ============================================================
    # CONSTRUIR QUERY (IGUAL QUE EN assets_list)
    # ============================================================
    
    assets = Asset.objects.all().prefetch_related(
        'layers',
        'layers__benchmark',
        Prefetch('scans', queryset=Scan.objects.order_by('-scan_date'))
    )
    
    # Aplicar filtros
    if status_filter:
        assets = assets.filter(status=status_filter)
    
    if search_query:
        assets = assets.filter(
            Q(hostname__icontains=search_query) | 
            Q(ip_address__icontains=search_query)
        )
    
    if environment_filter:
        assets = assets.filter(environment=environment_filter)
    
    if administrator_filter:
        assets = assets.filter(administrator=administrator_filter)
    
    if responsible_filter:
        assets = assets.filter(responsible=responsible_filter)
    
    if pci_filter:
        assets = assets.filter(pci_standard=(pci_filter == 'true'))
    
    if layer_type_filter:
        assets = assets.filter(layers__layer_type=layer_type_filter).distinct()
    
    # ============================================================
    # ENRIQUECER DATOS
    # ============================================================
    
    enriched_assets = []
    
    for asset in assets:
        asset_data = {
            'asset': asset,
            'layers_data': []
        }
        
        for layer in asset.layers.all():
            latest_scan = layer.scans.order_by('-scan_date').first()
            
            try:
                baseline = Baseline.objects.get(asset=asset, layer=layer)
                baseline_value = float(baseline.compliance_percentage)
            except Baseline.DoesNotExist:
                baseline_value = None
            
            layer_info = {
                'layer': layer,
                'latest_scan': latest_scan,
                'baseline': baseline_value,
                'compliance_change': None,
                'compliance_changed': False
            }
            
            if latest_scan and baseline_value is not None:
                current_compliance = float(latest_scan.compliance_percentage)
                layer_info['compliance_change'] = current_compliance - baseline_value
                layer_info['compliance_changed'] = abs(layer_info['compliance_change']) >= 1.0
            
            asset_data['layers_data'].append(layer_info)
        
        enriched_assets.append(asset_data)
    
    # ============================================================
    # APLICAR FILTROS AVANZADOS (IGUAL QUE EN LA VISTA)
    # ============================================================
    
    if scan_date_from or scan_date_to:
        filtered_assets = []
        for asset_data in enriched_assets:
            has_valid_scan = False
            for layer_data in asset_data['layers_data']:
                if layer_data['latest_scan']:
                    scan_date = layer_data['latest_scan'].scan_date
                    
                    if scan_date_from and scan_date < datetime.strptime(scan_date_from, '%Y-%m-%d').date():
                        continue
                    if scan_date_to and scan_date > datetime.strptime(scan_date_to, '%Y-%m-%d').date():
                        continue
                    
                    has_valid_scan = True
                    break
            
            if has_valid_scan:
                filtered_assets.append(asset_data)
        
        enriched_assets = filtered_assets
    
    if compliance_min or compliance_max:
        filtered_assets = []
        for asset_data in enriched_assets:
            has_valid_compliance = False
            for layer_data in asset_data['layers_data']:
                if layer_data['latest_scan']:
                    compliance = float(layer_data['latest_scan'].compliance_percentage)
                    
                    if compliance_min and compliance < float(compliance_min):
                        continue
                    if compliance_max and compliance > float(compliance_max):
                        continue
                    
                    has_valid_compliance = True
                    break
            
            if has_valid_compliance:
                filtered_assets.append(asset_data)
        
        enriched_assets = filtered_assets
    
    if show_only_changed:
        filtered_assets = []
        for asset_data in enriched_assets:
            has_changed = any(
                layer_data['compliance_changed'] 
                for layer_data in asset_data['layers_data']
            )
            if has_changed:
                filtered_assets.append(asset_data)
        
        enriched_assets = filtered_assets
    
    # ============================================================
    # CREAR ARCHIVO EXCEL
    # ============================================================
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Assets Compliance"
    
    # Estilos
    header_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    
    cell_alignment_left = Alignment(horizontal='left', vertical='center')
    cell_alignment_center = Alignment(horizontal='center', vertical='center')
    
    border = Border(
        left=Side(style='thin', color='D0D0D0'),
        right=Side(style='thin', color='D0D0D0'),
        top=Side(style='thin', color='D0D0D0'),
        bottom=Side(style='thin', color='D0D0D0')
    )
    
    # ============================================================
    # ENCABEZADOS
    # ============================================================
    
    headers = [
        'Hostname',
        'IP Address',
        'Estado',
        'Estándar PCI',
        'Ambiente',
        'Administrador',
        'Responsable',
        'Capa',
        'Tipo Software',
        'Benchmark',
        'Compliance (%)',
        'Línea Base (%)',
        'Cambio vs LB',
        'Passed',
        'Failed',
        'Warning',
        'Último Escaneo',
        'Total Políticas'
    ]
    
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = border
    
    # ============================================================
    # DATOS
    # ============================================================
    
    row_num = 2
    
    for asset_data in enriched_assets:
        asset = asset_data['asset']
        
        # Si no tiene capas, mostrar una fila con datos básicos
        if not asset_data['layers_data']:
            ws.cell(row=row_num, column=1, value=asset.hostname).alignment = cell_alignment_left
            ws.cell(row=row_num, column=2, value=asset.ip_address).alignment = cell_alignment_center
            ws.cell(row=row_num, column=3, value=asset.get_status_display()).alignment = cell_alignment_center
            ws.cell(row=row_num, column=4, value='Sí' if asset.pci_standard else 'No').alignment = cell_alignment_center
            ws.cell(row=row_num, column=5, value=asset.get_environment_display() if asset.environment else '-').alignment = cell_alignment_left
            ws.cell(row=row_num, column=6, value=asset.get_administrator_display() if asset.administrator else '-').alignment = cell_alignment_left
            ws.cell(row=row_num, column=7, value=asset.get_responsible_display() if asset.responsible else '-').alignment = cell_alignment_left
            
            # Aplicar bordes
            for col in range(1, len(headers) + 1):
                ws.cell(row=row_num, column=col).border = border
            
            row_num += 1
        else:
            # Una fila por cada capa
            for layer_data in asset_data['layers_data']:
                layer = layer_data['layer']
                latest_scan = layer_data['latest_scan']
                baseline = layer_data['baseline']
                compliance_change = layer_data['compliance_change']
                
                # Datos básicos del asset
                ws.cell(row=row_num, column=1, value=asset.hostname).alignment = cell_alignment_left
                ws.cell(row=row_num, column=2, value=asset.ip_address).alignment = cell_alignment_center
                ws.cell(row=row_num, column=3, value=asset.get_status_display()).alignment = cell_alignment_center
                ws.cell(row=row_num, column=4, value='Sí' if asset.pci_standard else 'No').alignment = cell_alignment_center
                ws.cell(row=row_num, column=5, value=asset.get_environment_display() if asset.environment else '-').alignment = cell_alignment_left
                ws.cell(row=row_num, column=6, value=asset.get_administrator_display() if asset.administrator else '-').alignment = cell_alignment_left
                ws.cell(row=row_num, column=7, value=asset.get_responsible_display() if asset.responsible else '-').alignment = cell_alignment_left
                
                # Datos de la capa
                ws.cell(row=row_num, column=8, value=layer.get_layer_type_display()).alignment = cell_alignment_center
                ws.cell(row=row_num, column=9, value=layer.get_software_type_display() if layer.software_type else '-').alignment = cell_alignment_left
                ws.cell(row=row_num, column=10, value=str(layer.benchmark)).alignment = cell_alignment_left
                
                # Datos del escaneo
                if latest_scan:
                    compliance_cell = ws.cell(row=row_num, column=11, value=float(latest_scan.compliance_percentage))
                    compliance_cell.number_format = '0.00'
                    compliance_cell.alignment = cell_alignment_center
                    
                    # Colorear según compliance
                    if latest_scan.compliance_percentage >= 90:
                        compliance_cell.fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
                        compliance_cell.font = Font(color='006100', bold=True)
                    elif latest_scan.compliance_percentage >= 70:
                        compliance_cell.fill = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')
                        compliance_cell.font = Font(color='9C6500', bold=True)
                    else:
                        compliance_cell.fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
                        compliance_cell.font = Font(color='9C0006', bold=True)
                    
                    # Línea base
                    if baseline is not None:
                        baseline_cell = ws.cell(row=row_num, column=12, value=baseline)
                        baseline_cell.number_format = '0.00'
                        baseline_cell.alignment = cell_alignment_center
                    else:
                        ws.cell(row=row_num, column=12, value='-').alignment = cell_alignment_center
                    
                    # Cambio vs línea base
                    if compliance_change is not None:
                        change_cell = ws.cell(row=row_num, column=13, value=compliance_change)
                        change_cell.number_format = '+0.00;-0.00;0.00'
                        change_cell.alignment = cell_alignment_center
                        
                        # Colorear según cambio
                        if compliance_change > 0:
                            change_cell.font = Font(color='006100', bold=True)
                        elif compliance_change < 0:
                            change_cell.font = Font(color='9C0006', bold=True)
                    else:
                        ws.cell(row=row_num, column=13, value='-').alignment = cell_alignment_center
                    
                    # Métricas
                    ws.cell(row=row_num, column=14, value=latest_scan.passed_count).alignment = cell_alignment_center
                    ws.cell(row=row_num, column=15, value=latest_scan.failed_count).alignment = cell_alignment_center
                    ws.cell(row=row_num, column=16, value=latest_scan.warning_count).alignment = cell_alignment_center
                    ws.cell(row=row_num, column=17, value=latest_scan.scan_date.strftime('%Y-%m-%d')).alignment = cell_alignment_center
                    ws.cell(row=row_num, column=18, value=latest_scan.total_policies).alignment = cell_alignment_center
                else:
                    # Sin escaneo
                    for col in range(11, 19):
                        ws.cell(row=row_num, column=col, value='-').alignment = cell_alignment_center
                
                # Aplicar bordes
                for col in range(1, len(headers) + 1):
                    ws.cell(row=row_num, column=col).border = border
                
                row_num += 1
    
    # ============================================================
    # AJUSTAR ANCHOS DE COLUMNA
    # ============================================================
    
    column_widths = {
        1: 30,  # Hostname
        2: 15,  # IP
        3: 12,  # Estado
        4: 12,  # PCI
        5: 20,  # Ambiente
        6: 25,  # Administrador
        7: 25,  # Responsable
        8: 15,  # Capa
        9: 25,  # Tipo Software
        10: 35, # Benchmark
        11: 14, # Compliance
        12: 14, # Línea Base
        13: 14, # Cambio
        14: 10, # Passed
        15: 10, # Failed
        16: 10, # Warning
        17: 15, # Fecha
        18: 14, # Total
    }
    
    for col_num, width in column_widths.items():
        ws.column_dimensions[get_column_letter(col_num)].width = width
    
    # Congelar primera fila
    ws.freeze_panes = 'A2'
    
    # ============================================================
    # AGREGAR HOJA DE RESUMEN
    # ============================================================
    
    # ws_summary = wb.create_sheet("Resumen")
    
    # # Título
    # ws_summary['A1'] = 'RESUMEN DE EXPORTACIÓN'
    # ws_summary['A1'].font = Font(size=14, bold=True, color='4472C4')
    # ws_summary.merge_cells('A1:B1')
    
    # # Información
    # summary_data = [
    #     ['Fecha de Exportación:', datetime.now().strftime('%Y-%m-%d %H:%M:%S')],
    #     ['Usuario:', request.user.username],
    #     ['Total de Assets:', len(enriched_assets)],
    #     [''],
    #     ['FILTROS APLICADOS:', ''],
    #     ['Estado:', status_filter if status_filter else 'Todos'],
    #     ['Búsqueda:', search_query if search_query else 'Sin filtro'],
    #     ['Ambiente:', environment_filter if environment_filter else 'Todos'],
    #     ['Administrador:', administrator_filter if administrator_filter else 'Todos'],
    #     ['Responsable:', responsible_filter if responsible_filter else 'Todos'],
    #     ['Tipo de Capa:', layer_type_filter if layer_type_filter else 'Todas'],
    #     ['Estándar PCI:', pci_filter if pci_filter else 'Todos'],
    #     ['Compliance Mínimo:', f'{compliance_min}%' if compliance_min else 'Sin límite'],
    #     ['Compliance Máximo:', f'{compliance_max}%' if compliance_max else 'Sin límite'],
    #     ['Fecha Desde:', scan_date_from if scan_date_from else 'Sin límite'],
    #     ['Fecha Hasta:', scan_date_to if scan_date_to else 'Sin límite'],
    #     ['Solo con cambios:', 'Sí' if show_only_changed else 'No'],
    # ]
    
    # for row_num, (label, value) in enumerate(summary_data, start=2):
    #     ws_summary.cell(row=row_num, column=1, value=label).font = Font(bold=True)
    #     ws_summary.cell(row=row_num, column=2, value=value)
    
    # ws_summary.column_dimensions['A'].width = 25
    # ws_summary.column_dimensions['B'].width = 40
    
    # ============================================================
    # PREPARAR RESPUESTA HTTP
    # ============================================================
    
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    
    filename = f'Assets_Compliance_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    
    wb.save(response)
    
    return response